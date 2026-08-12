import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

APP_DIR = Path(__file__).resolve().parents[1]
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from detector import PIIDetector
from docx_io import DocxIO
from app import app


class DetectorTests(unittest.TestCase):
    def test_does_not_treat_official_heading_as_person(self):
        text = (
            'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n'
            'VIỆN KIỂM SÁT NHÂN DÂN TỈNH'
        )
        names = [item for item in PIIDetector(text).detect() if item['type'] == 'Tên']
        self.assertEqual([], names)

    def test_detects_distinct_people_with_similar_names(self):
        text = (
            'Ông Nguyễn Văn Anh và bà Trần Thị Bình có mặt. '
            'Nguyễn Văn An vắng mặt.'
        )
        names = {
            item['value'] for item in PIIDetector(text).detect()
            if item['type'] == 'Tên'
        }
        self.assertEqual({'Nguyễn Văn Anh', 'Trần Thị Bình', 'Nguyễn Văn An'}, names)

    def test_address_is_not_reported_again_as_a_person(self):
        text = (
            'Bị cáo Nguyễn Văn An, trú tại số 12, đường Lê Lợi, '
            'phường A, thành phố B.'
        )
        entities = PIIDetector(text).detect()
        names = {item['value'] for item in entities if item['type'] == 'Tên'}
        addresses = {item['value'] for item in entities if item['type'] == 'Địa chỉ'}
        self.assertEqual({'Nguyễn Văn An'}, names)
        self.assertEqual({'số 12, đường Lê Lợi, phường A, thành phố B'}, addresses)


class DocxReplacementTests(unittest.TestCase):
    def test_replaces_text_split_across_runs_and_in_header_table(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'source.docx'
            output = Path(temp_dir) / 'output.docx'

            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run('Ông Nguyễn ').bold = True
            paragraph.add_run('Văn An').italic = True
            paragraph.add_run(' có mặt.')
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = 'Nguyễn Văn An ký tên.'
            document.sections[0].header.paragraphs[0].text = 'Hồ sơ Nguyễn Văn An'
            document.save(source)

            docx = DocxIO(source)
            count = docx.replace_text('Nguyễn Văn An', 'Tên-A')
            docx.save(output)

            result = DocxIO(output).get_text()
            self.assertEqual(3, count)
            self.assertNotIn('Nguyễn Văn An', result)
            self.assertEqual(3, result.count('Tên-A'))

    def test_does_not_replace_inside_a_longer_word(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'source.docx'
            document = Document()
            document.add_paragraph('Mã AN123 và mã AN là hai giá trị.')
            document.save(source)

            docx = DocxIO(source)
            count = docx.replace_text('AN', '[Ẩn]')
            self.assertEqual(1, count)
            self.assertIn('AN123', docx.get_text())
            self.assertIn('mã [Ẩn]', docx.get_text())


class ApiTests(unittest.TestCase):
    def test_manual_replacement_is_exported(self):
        source = BytesIO()
        document = Document()
        document.add_paragraph('Ông Nguyễn Văn An quản lý Mã hồ sơ nội bộ XYZ.')
        document.save(source)
        source.seek(0)

        app.config['TESTING'] = True
        with app.test_client() as client:
            upload = client.post(
                '/upload',
                data={'file': (source, 'sample.docx')},
                content_type='multipart/form-data',
            )
            self.assertEqual(200, upload.status_code)
            detected = {item['value'] for item in upload.get_json()['entities']}
            self.assertIn('Nguyễn Văn An', detected)

            exported = client.post('/export', json={'replacements': [
                {'old': 'Nguyễn Văn An', 'new': 'Tên-A'},
                {'old': 'Mã hồ sơ nội bộ XYZ', 'new': '[Thông tin nội bộ]'},
            ]})
            self.assertEqual(200, exported.status_code)
            exported_bytes = exported.data
            exported.close()
            result = Document(BytesIO(exported_bytes))
            text = '\n'.join(paragraph.text for paragraph in result.paragraphs)
            self.assertIn('Tên-A', text)
            self.assertIn('[Thông tin nội bộ]', text)
            self.assertNotIn('Nguyễn Văn An', text)
            self.assertNotIn('Mã hồ sơ nội bộ XYZ', text)


if __name__ == '__main__':
    unittest.main()
