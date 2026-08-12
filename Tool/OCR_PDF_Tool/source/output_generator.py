"""
Output Generator Module - Xuất kết quả OCR ra các định dạng
Hỗ trợ: .md (Markdown), .txt (Text), .docx (Word)
"""

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from config import DOCX_FONT_NAME, DOCX_FONT_SIZE


class OutputGenerator:
    """Tạo file output từ kết quả OCR."""

    def __init__(self, file_name: str, output_dir: str):
        """
        Khởi tạo Output Generator.

        Args:
            file_name: Tên file gốc (không có đuôi mở rộng)
            output_dir: Thư mục lưu file output
        """
        self.file_name = file_name
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate(self, pages_text: list, output_format: str) -> str:
        """
        Tạo file output theo định dạng yêu cầu.

        Args:
            pages_text: Danh sách text của từng trang (list[str])
            output_format: 'md', 'txt', hoặc 'docx'

        Returns:
            Đường dẫn đến file đã tạo
        """
        generators = {
            "md": self._generate_markdown,
            "txt": self._generate_text,
            "docx": self._generate_docx,
        }

        generator = generators.get(output_format)
        if generator is None:
            raise ValueError(
                f"Định dạng không hỗ trợ: {output_format}. "
                f"Hỗ trợ: {', '.join(generators.keys())}"
            )

        return generator(pages_text)

    def _get_output_path(self, extension: str) -> str:
        """Tạo đường dẫn file output."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{self.file_name}_OCR_{timestamp}.{extension}"
        return os.path.join(self.output_dir, filename)

    def _generate_markdown(self, pages_text: list) -> str:
        """
        Xuất kết quả ra file Markdown (.md).

        Format:
            # Kết quả OCR - {Tên File}

            ## Trang 1
            Nội dung...

            ## Trang 2
            Nội dung...
        """
        output_path = self._get_output_path("md")

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"# Kết quả OCR - {self.file_name}\n\n")
            f.write(f"> Ngày tạo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n\n")
            f.write("---\n\n")

            for i, text in enumerate(pages_text, 1):
                f.write(f"## Trang {i}\n\n")
                if text.strip():
                    f.write(f"{text.strip()}\n")
                else:
                    f.write("*Không nhận dạng được văn bản*\n")
                f.write("\n\n")

            f.write("---\n\n")
            f.write(f"*Tổng số trang: {len(pages_text)}*\n")

        return output_path

    def _generate_text(self, pages_text: list) -> str:
        """
        Xuất kết quả ra file Text (.txt).

        Format:
            === Trang 1 ===
            Nội dung...

            === Trang 2 ===
            Nội dung...
        """
        output_path = self._get_output_path("txt")

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"KET QUA OCR - {self.file_name}\n")
            f.write(f"Ngay tao: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
            f.write("=" * 60 + "\n\n")

            for i, text in enumerate(pages_text, 1):
                f.write(f"=== Trang {i} ===\n\n")
                if text.strip():
                    f.write(f"{text.strip()}\n")
                else:
                    f.write("[Khong nhan dang duoc van ban]\n")
                f.write("\n" + "-" * 40 + "\n\n")

            f.write(f"Tong so trang: {len(pages_text)}\n")

        return output_path

    def _generate_docx(self, pages_text: list) -> str:
        """
        Xuất kết quả ra file Word (.docx).

        Format:
            - Tiêu đề: "Kết quả OCR - {Tên File}"
            - Mỗi trang là 1 section với heading "Trang X"
            - Font: Times New Roman, size 13
            - Page break giữa các trang
        """
        output_path = self._get_output_path("docx")
        doc = Document()

        # ── Cấu hình style mặc định ──────────────────────────
        style = doc.styles["Normal"]
        font = style.font
        font.name = DOCX_FONT_NAME
        font.size = Pt(DOCX_FONT_SIZE)

        # Đảm bảo font Unicode (tiếng Việt) cho cả chữ thường
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = __import__("lxml.etree", fromlist=["etree"]).SubElement(rPr, qn("w:rFonts"))
        rFonts.set(qn("w:eastAsia"), DOCX_FONT_NAME)
        rFonts.set(qn("w:cs"), DOCX_FONT_NAME)

        # ── Tiêu đề chính ────────────────────────────────────
        title = doc.add_heading(f"Kết quả OCR - {self.file_name}", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.name = DOCX_FONT_NAME
            run.font.size = Pt(18)

        # Ngày tạo
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_run = date_para.add_run(
            f"Ngày tạo: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        )
        date_run.font.size = Pt(10)
        date_run.font.italic = True

        doc.add_paragraph()  # Dòng trống

        # ── Nội dung từng trang ──────────────────────────────
        for i, text in enumerate(pages_text, 1):
            # Heading trang
            heading = doc.add_heading(f"Trang {i}", level=2)
            for run in heading.runs:
                run.font.name = DOCX_FONT_NAME

            # Nội dung
            if text.strip():
                para = doc.add_paragraph(text.strip())
                para.style.font.name = DOCX_FONT_NAME
                para.style.font.size = Pt(DOCX_FONT_SIZE)
            else:
                para = doc.add_paragraph("[Không nhận dạng được văn bản]")
                para.style.font.italic = True

            # Page break (trừ trang cuối)
            if i < len(pages_text):
                doc.add_page_break()

        # ── Footer ───────────────────────────────────────────
        doc.add_paragraph()
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_para.add_run(
            f"Tổng số trang: {len(pages_text)}"
        )
        footer_run.font.size = Pt(10)
        footer_run.font.italic = True

        # Lưu file
        doc.save(output_path)

        return output_path