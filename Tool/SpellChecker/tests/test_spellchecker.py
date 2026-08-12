import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx_processor import DocxReview
from engine import VietnameseSpellChecker
from app import app


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.checker = VietnameseSpellChecker()

    def test_known_legal_phrase_and_spacing(self):
        text = "Viện kiểm xát nhân dân ban hành  quyết định khởi tộ bị can."
        issues = self.checker.check_text(text, "p1", "Thân văn bản")
        pairs = {(item.original.casefold(), item.suggestion.casefold()) for item in issues}
        self.assertTrue(
            ("kiểm xát", "kiểm sát") in pairs
            or ("viện kiểm xát", "viện kiểm sát") in pairs
        )
        self.assertIn(("khởi tộ", "khởi tố"), pairs)
        self.assertTrue(any(item.kind == "Trình bày" for item in issues))

    def test_protected_identifiers_are_not_reported(self):
        text = "VKSND ban hành Quyết định số 12/QĐ-VKS; bị can Nguyễn Khắc Tú."
        issues = self.checker.check_text(text)
        originals = {item.original for item in issues}
        self.assertNotIn("VKSND", originals)
        self.assertNotIn("Khắc", originals)

    def test_combining_unicode_is_accepted(self):
        decomposed = "quyết định"
        issues = self.checker.check_text(decomposed)
        self.assertFalse(any(item.kind == "Unicode" for item in issues))
        self.assertFalse(any(item.original in {"quyết", "định"} for item in issues))

    def test_modern_tone_placement_is_accepted(self):
        text = "Tòa án ghi nhận các bên thỏa thuận hòa giải và tùy nghi áp dụng."
        issues = self.checker.check_text(text)
        originals = {item.original.casefold() for item in issues}
        self.assertTrue({"tòa", "thỏa", "hòa", "tùy"}.isdisjoint(originals))

    def test_measurement_units_are_not_reported(self):
        text = (
            "Khoảng cách 12 cm, 8 mm, 2 km; khối lượng 250 mg, 3 kg; "
            "áp suất 2 MPa; kích thước 10-12 cm và đoạn còn lại dài 5cm."
        )
        issues = self.checker.check_text(text)
        originals = {item.original.casefold() for item in issues}
        self.assertTrue({"cm", "mm", "km", "mg", "kg", "mpa"}.isdisjoint(originals))

    def test_measurement_like_token_without_number_is_still_checked(self):
        issues = self.checker.check_text("Văn bản ghi cm trong một cụm từ nhưng không có số liệu.")
        self.assertTrue(any(item.original.casefold() == "cm" for item in issues))


class DocxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.checker = VietnameseSpellChecker()

    def make_doc(self, path):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Viện kiểm ").bold = True
        paragraph.add_run("xát nhân dân ban hành  quyết định khởi tộ.")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Bị cao Nguyễn Văn A."
        doc.sections[0].header.paragraphs[0].text = "VKSND KHU VỰC"
        doc.save(path)

    def test_analyze_mark_and_correct_docx(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            marked = Path(folder) / "marked.docx"
            corrected = Path(folder) / "corrected.docx"
            self.make_doc(source)

            review = DocxReview(source, self.checker)
            issues, count = review.analyze()
            self.assertGreater(count, 5)
            self.assertTrue(any("kiểm xát" in item.original.casefold() for item in issues))
            self.assertTrue(any(item.original.casefold() == "bị cao" for item in issues))

            review.save_marked(marked)
            self.assertTrue(marked.exists())
            marked_doc = Document(marked)
            self.assertIn("kiểm xát", "\n".join(p.text for p in marked_doc.paragraphs))
            self.assertEqual(len(marked_doc.comments), 0)
            highlighted = [
                run.font.highlight_color
                for paragraph in marked_doc.paragraphs
                for run in paragraph.runs
                if run.font.highlight_color is not None
            ]
            self.assertGreater(len(highlighted), 0)
            self.assertTrue(all(color == WD_COLOR_INDEX.YELLOW for color in highlighted))

            accepted = {item.id for item in issues if item.severity == "error"}
            fresh = DocxReview(source, self.checker)
            fresh.analyze()
            fresh.save_corrected(corrected, accepted)
            corrected_doc = Document(corrected)
            body = "\n".join(p.text for p in corrected_doc.paragraphs)
            self.assertIn("kiểm sát", body)
            self.assertIn("khởi tố", body)

    def test_http_analyze_and_export(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "api.docx"
            self.make_doc(source)
            client = app.test_client()
            with source.open("rb") as stream:
                response = client.post(
                    "/api/analyze",
                    data={"file": (stream, "api.docx")},
                    content_type="multipart/form-data",
                )
            self.assertEqual(response.status_code, 200)
            payload = response.get_json()
            self.assertGreater(len(payload["issues"]), 0)
            accepted = [item["id"] for item in payload["issues"] if item["severity"] == "error"]
            exported = client.post(
                "/api/export",
                json={"job": payload["job"], "mode": "both", "accepted": accepted},
            )
            self.assertEqual(exported.status_code, 200)
            with zipfile.ZipFile(io.BytesIO(exported.data)) as archive:
                self.assertEqual(len(archive.namelist()), 2)
            exported.close()


if __name__ == "__main__":
    unittest.main()
