"""Đọc, đánh dấu và sửa DOCX trong khi giữ tối đa định dạng gốc."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run

from engine import Issue, VietnameseSpellChecker


@dataclass
class ParagraphRef:
    key: str
    location: str
    paragraph: object


def iter_paragraphs(document):
    """Duyệt thân bài, bảng lồng nhau, đầu trang và chân trang; tránh lặp XML."""
    seen = set()
    counter = 0

    def emit(paragraphs, label):
        nonlocal counter
        for paragraph in paragraphs:
            marker = paragraph._p
            if marker in seen:
                continue
            seen.add(marker)
            counter += 1
            yield ParagraphRef(f"p{counter}", label, paragraph)

    yield from emit(document.paragraphs, "Thân văn bản")

    def walk_tables(tables, prefix):
        for table_index, table in enumerate(tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                for cell_index, cell in enumerate(row.cells, 1):
                    label = f"{prefix}Bảng {table_index}, hàng {row_index}, cột {cell_index}"
                    yield from emit(cell.paragraphs, label)
                    yield from walk_tables(cell.tables, label + " / ")

    yield from walk_tables(document.tables, "")
    for section_index, section in enumerate(document.sections, 1):
        yield from emit(section.header.paragraphs, f"Đầu trang, phần {section_index}")
        yield from emit(section.footer.paragraphs, f"Chân trang, phần {section_index}")


class DocxReview:
    def __init__(self, path: str | Path, checker: VietnameseSpellChecker):
        self.path = Path(path)
        self.document = Document(str(self.path))
        self.checker = checker
        self.refs = {ref.key: ref for ref in iter_paragraphs(self.document)}
        self.issues: list[Issue] = []

    def analyze(self):
        issues = []
        word_count = 0
        for ref in self.refs.values():
            text = ref.paragraph.text
            word_count += len(text.split())
            if text.strip():
                issues.extend(self.checker.check_text(text, ref.key, ref.location))
        # ID ổn định và duy nhất trong toàn tài liệu.
        for index, issue in enumerate(issues, 1):
            issue.id = f"issue-{index}"
        self.issues = issues
        return issues, word_count

    @staticmethod
    def _run_text(run_element):
        return Run(run_element, None).text

    @staticmethod
    def _new_run_like(source_run, text):
        element = OxmlElement("w:r")
        if source_run._r.rPr is not None:
            element.append(deepcopy(source_run._r.rPr))
        node = OxmlElement("w:t")
        if text.startswith(" ") or text.endswith(" "):
            node.set(qn("xml:space"), "preserve")
        node.text = text
        element.append(node)
        return element

    def _split_for_range(self, paragraph, start, end):
        """Tách các run tại biên và trả về run nằm đúng trong khoảng ký tự."""
        selected = []
        position = 0
        for run in list(paragraph.runs):
            text = run.text
            run_start, run_end = position, position + len(text)
            position = run_end
            if not text or end <= run_start or start >= run_end:
                continue
            local_start = max(0, start - run_start)
            local_end = min(len(text), end - run_start)
            before, middle, after = text[:local_start], text[local_start:local_end], text[local_end:]
            additions = []
            if before:
                additions.append(self._new_run_like(run, before))
            middle_element = self._new_run_like(run, middle)
            additions.append(middle_element)
            if after:
                additions.append(self._new_run_like(run, after))
            for element in additions:
                run._r.addprevious(element)
            run._r.getparent().remove(run._r)
            selected.append(Run(middle_element, paragraph))
        return selected

    def save_marked(self, output_path: str | Path):
        if not self.issues:
            self.analyze()
        grouped = {}
        for issue in self.issues:
            grouped.setdefault(issue.paragraph_id, []).append(issue)
        # Đi từ phải sang trái để tọa độ ký tự không thay đổi.
        for paragraph_id, issues in grouped.items():
            paragraph = self.refs[paragraph_id].paragraph
            for issue in sorted(issues, key=lambda item: item.start, reverse=True):
                runs = self._split_for_range(paragraph, issue.start, issue.end)
                if not runs:
                    continue
                for run in runs:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        self.document.save(str(output_path))

    def save_corrected(self, output_path: str | Path, accepted_ids: set[str]):
        if not self.issues:
            self.analyze()
        grouped = {}
        for issue in self.issues:
            if issue.id in accepted_ids and issue.suggestion:
                grouped.setdefault(issue.paragraph_id, []).append(issue)
        for paragraph_id, issues in grouped.items():
            paragraph = self.refs[paragraph_id].paragraph
            for issue in sorted(issues, key=lambda item: item.start, reverse=True):
                runs = self._split_for_range(paragraph, issue.start, issue.end)
                if not runs:
                    continue
                runs[0].text = issue.suggestion
                for run in runs[1:]:
                    run._r.getparent().remove(run._r)
        self.document.save(str(output_path))
