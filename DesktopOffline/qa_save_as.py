"""Kiểm thử xuất DOCX qua API native và đường dẫn Save As của Windows."""

from __future__ import annotations

import os
import tempfile
import threading
from pathlib import Path

from docx import Document
from werkzeug.serving import make_server

import main as desktop


STAGE = Path(__file__).resolve().parent / "stage" / "payload"
desktop.resource_root = lambda: STAGE
os.environ["BCKS_DESKTOP_OFFLINE"] = "1"


with tempfile.TemporaryDirectory() as temporary:
    temp = Path(temporary)
    source = temp / "nguon.docx"
    target = temp / "ket_qua.docx"
    document = Document()
    document.add_paragraph("Nguyễn Văn A là người cần ẩn danh.")
    document.save(source)

    tool = desktop.TOOLS["anonymizer"]
    module = desktop.load_source_module(STAGE / tool["source"])
    app = module.app
    app.template_folder = str(STAGE / tool["template_folder"])
    app.static_folder = str(STAGE / tool["static_folder"])
    app.config["LAST_DOC"] = str(source)

    server = make_server("127.0.0.1", 0, app, threaded=True)
    port = server.server_port
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        api = desktop.ToolApi(f"http://127.0.0.1:{port}")
        dialog_calls = []
        def fake_save_dialog(filename, file_types):
            dialog_calls.append((filename, file_types))
            return {"ok": True, "cancelled": False, "path": str(target)}
        api._save_path_dialog = fake_save_dialog
        result = api.save_anonymized([{"old": "Nguyễn Văn A", "new": "[Ẩn]"}])
    finally:
        server.shutdown()
        thread.join(timeout=5)

    if not result.get("ok") or result.get("cancelled"):
        raise SystemExit(f"Save As thất bại: {result}")
    if not target.exists() or target.stat().st_size < 1000:
        raise SystemExit("File Save As không tồn tại hoặc bị rỗng.")
    output = "\n".join(paragraph.text for paragraph in Document(target).paragraphs)
    if "[Ẩn]" not in output or "Nguyễn Văn A" in output:
        raise SystemExit(f"Nội dung file xuất không đúng: {output}")
    if len(dialog_calls) != 1:
        raise SystemExit("Hộp thoại Save As không được gọi đúng một lần.")
    print(f"OK Save As: {target.name}, {target.stat().st_size} bytes, nội dung đã ẩn danh")
